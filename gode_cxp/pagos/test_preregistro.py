"""Pre-registro de cuentas de proveedor en BancaNet: filas de la plantilla, respuesta del banco y
candado del lote.

El banco exige dar de alta la cuenta antes de poder pagarle al proveedor. La app NO genera el
archivo de alta (sólo hay una muestra, de cuentas Banamex de cheques de personas físicas): produce
las filas para pegar en la plantilla con macros del banco y lee la respuesta.
"""
import io
import json
import os
import re
import unittest
from datetime import date

import frappe
import openpyxl
from frappe.tests.utils import FrappeTestCase

from gode_cxp.cfdi import ejemplos
from gode_cxp.facturas import pruebas_comun
from gode_cxp.facturas.pruebas_comun import cuenta_verificada, factura_aprobada, usuario
from gode_cxp.pagos import api, preregistro
from gode_cxp.pagos.lotes import crear_lotes
from gode_cxp.setup.produccion import configurar_pagos

CLABE_12 = "072180007090045065"          # Banorte: interbancario, naturaleza 12
OTRA_CLABE_12 = "072180007090045997"     # otra CLABE válida del mismo banco
CLABE_06 = "002180012345678906"          # Banamex: naturaleza 06 (sucursal + cuenta)
CLABE_DESCONOCIDA = "998180007090045068"  # clave 998: no está en el catálogo del banco
SUCURSAL, CUENTA = "7005", "7479513"
TESORERIA, REVISOR = "prueba.tesoreria@cxp.local", "prueba.revisor@cxp.local"
FECHA = date(2026, 9, 17)
# La muestra REAL de una respuesta de alta vive fuera del repo (trae datos personales de empleados):
# bench-pruebas.sh monta /home/sergio/service-env/tef-muestras aquí, de sólo lectura.
MUESTRA_REAL = "/home/frappe/tef-muestras/abcMasivoTERINT170920260001815117770000000000020001.txt"


def registro(banco, tipo, cuenta20, beneficiario, codigo, mensaje):
    """Un registro de detalle (325) del archivo de respuesta, armado con las posiciones A MANO y no
    con las constantes del módulo: si el lector se equivoca de posición, la prueba tiene que notarlo.

    'A' (1) + banco (2-5) + tipo de cuenta (6-7) + un número fijo del contrato (8-19) +
    cuenta a 20 (20-39) + beneficiario (40-94) + … + código (266-269) + mensaje (270-325).
    """
    linea = "A" + banco + tipo + "0" * 12 + cuenta20 + beneficiario.ljust(55)
    assert len(linea) == 94, len(linea)
    linea = linea.ljust(265) + codigo + mensaje.ljust(56)
    assert len(linea) == 325, len(linea)
    return linea


def archivo_respuesta(registros):
    """Encabezado de 115 + los registros, latin-1, CRLF ENTRE líneas y sin CRLF al final (así llega
    el archivo real del banco)."""
    encabezado = ("0001" + "17/09/2026" + "05:26" + "000181511777").ljust(115)
    assert len(encabezado) == 115
    return "\r\n".join([encabezado] + registros).encode("latin-1")


def subir(contenido, nombre="abcMasivo-prueba.txt"):
    """Deja el archivo como File privado suelto y devuelve su file_url (es lo que sube Tesorería)."""
    return frappe.get_doc({"doctype": "File", "file_name": nombre, "content": contenido,
                           "is_private": 1}).insert(ignore_permissions=True).file_url


class TestPreregistro(FrappeTestCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        pruebas_comun.preparar_sitio_pruebas()
        banco = frappe.db.get_value("Account", {"company": pruebas_comun.EMPRESA, "account_type": "Bank", "is_group": 0}, "name")
        configurar_pagos(pruebas_comun.EMPRESA, "000181511777", "7007", "8382129",
                         "GASTRONOMICA DE ESPECIALIDADES GODE", "pago gode", banco, dry_run=False)

    def setUp(self):
        frappe.set_user("Administrator")
        pruebas_comun.limpiar()
        usuario(TESORERIA, "CxP Tesoreria"); usuario(REVISOR, "CxP Revisor")
        # Proveedor MORAL (AVICOLA DEL CARMEN SA DE CV) con cuenta interbancaria, y proveedor FÍSICO
        # (BRUNO RICARDO HERNANDEZ SILVA) con cuenta Banamex: las dos formas de fila de la plantilla.
        self.moral = factura_aprobada(ejemplos.INGRESO_40)
        self.fisica = factura_aprobada(ejemplos.INGRESO_33_RETENCIONES)
        self.cta_12 = cuenta_verificada(self.moral.supplier, CLABE_12, registrada=False)
        self.cta_06 = cuenta_verificada(self.fisica.supplier, CLABE_06, banco="Banamex",
                                        sucursal=SUCURSAL, cuenta=CUENTA, registrada=False)

    def tearDown(self):
        frappe.set_user("Administrator")

    # --- las filas para la plantilla del banco -----------------------------------------------------

    def test_fila_de_una_cuenta_interbancaria_de_persona_moral(self):
        (fila,) = preregistro.filas_para_plantilla([self.cta_12])
        self.assertEqual(fila, {
            "MOVIMIENTO": "ALTA",
            "BANCO": "BANORTE",
            "SUCURSAL": "",
            "TIPO DE CUENTA": "CLABE INTERBANCARIA",
            "NUMERO DE CUENTA": CLABE_12,
            "TIPO DE PERSONA": "PERSONA MORAL",
            "PERIODO": "DIARIO",
            "IMPORTE MAXIMO": 500000,
            "BENEFICIARIO": "AVICOLA,DEL CARMEN SA DE CV/",
            "ALIAS": "AVICOLA DEL CARMEN S",
            "RFC": "AVI900101AB1",
            "CELULAR": "",
            "EMAIL": "",
        })

    def test_fila_de_una_cuenta_banamex_de_persona_fisica(self):
        (fila,) = preregistro.filas_para_plantilla([self.cta_06])
        self.assertEqual(fila, {
            "MOVIMIENTO": "ALTA",
            "BANCO": "BANAMEX",
            "SUCURSAL": SUCURSAL,
            "TIPO DE CUENTA": "CHEQUES",
            "NUMERO DE CUENTA": CUENTA,
            "TIPO DE PERSONA": "PERSONA FISICA",
            "PERIODO": "DIARIO",
            "IMPORTE MAXIMO": 500000,
            "BENEFICIARIO": "BRUNO RICARDO,HERNANDEZ/SILVA",
            "ALIAS": "BRUNO RICARDO HERNAN",
            "RFC": "HESB850101AB1",
            "CELULAR": "",
            "EMAIL": "",
        })

    def test_el_importe_maximo_de_la_cuenta_le_gana_al_de_la_configuracion(self):
        frappe.db.set_value("Bank Account", self.cta_12, "importe_maximo_banco", 12345)
        (fila,) = preregistro.filas_para_plantilla([self.cta_12])
        self.assertEqual(fila["IMPORTE MAXIMO"], 12345)

    def test_el_banco_sale_de_los_tres_primeros_digitos_de_la_clabe(self):
        """Y si la clave no está en el catálogo, va el código de 3 dígitos: el banco rechaza la fila
        y se ve por qué, en vez de mandar un hueco."""
        self.assertEqual(preregistro.banco_de_clabe(CLABE_06), "BANAMEX")
        self.assertEqual(preregistro.banco_de_clabe(CLABE_12), "BANORTE")
        self.assertEqual(preregistro.banco_de_clabe("012180007090045067"), "BBVA")
        self.assertEqual(preregistro.banco_de_clabe(CLABE_DESCONOCIDA), "998")
        self.assertEqual(preregistro.banco_de_clabe(""), "")

    def test_solo_salen_las_cuentas_verificadas_que_faltan_por_registrar(self):
        self.assertEqual({f["NUMERO DE CUENTA"] for f in preregistro.filas_para_plantilla()},
                         {CLABE_12, CUENTA})
        frappe.db.set_value("Bank Account", self.cta_12, "estado_preregistro", "Registrada")
        self.assertEqual({f["NUMERO DE CUENTA"] for f in preregistro.filas_para_plantilla()}, {CUENTA})
        # Una cuenta que Tesorería no ha verificado no se manda al banco ni pidiéndola por su nombre.
        frappe.db.set_value("Bank Account", self.cta_06, "verificada", 0)
        self.assertEqual(preregistro.filas_para_plantilla(), [])
        with self.assertRaises(frappe.ValidationError):
            preregistro.filas_para_plantilla([self.cta_06])

    # --- la descarga del XLSX ----------------------------------------------------------------------

    def test_descargar_arma_el_xlsx_y_marca_las_cuentas_enviadas(self):
        frappe.set_user(TESORERIA)
        r = api.descargar_preregistro()
        self.assertEqual(r["cuentas"], 2)
        self.assertTrue(re.fullmatch(r"preregistro-\d{8}-\d{4}\.xlsx", r["nombre"]), r["nombre"])
        contenido = frappe.get_doc("File", {"file_url": r["file_url"]}).get_content()
        libro = openpyxl.load_workbook(io.BytesIO(contenido))
        self.assertEqual(libro.sheetnames, ["LOIncorOP"])
        hoja = libro["LOIncorOP"]
        self.assertEqual([c.value for c in hoja[1]], list(preregistro.COLUMNAS))
        self.assertEqual(hoja.max_row, 3)          # encabezados + 2 cuentas
        self.assertEqual(hoja.max_column, len(preregistro.COLUMNAS))
        for cuenta in (self.cta_12, self.cta_06):
            datos = frappe.db.get_value("Bank Account", cuenta,
                                        ["estado_preregistro", "preregistro_enviado_el"], as_dict=True)
            self.assertEqual(datos.estado_preregistro, "Enviada al banco")
            self.assertTrue(datos.preregistro_enviado_el)
        # Ya enviadas, no hay nada que descargar otra vez.
        with self.assertRaises(frappe.ValidationError):
            api.descargar_preregistro()

    def test_descargar_es_de_tesoreria(self):
        frappe.set_user(REVISOR)
        with self.assertRaises(frappe.PermissionError):
            api.descargar_preregistro()

    # --- la respuesta del banco -------------------------------------------------------------------

    def test_leer_la_respuesta_del_banco(self):
        datos = archivo_respuesta([
            registro("0002", "01", "0" * 9 + SUCURSAL + CUENTA, "BRUNO RICARDO,HERNANDEZ/SILVA", "0000", "ALTA APLICADA"),
            registro("0072", "40", "00" + CLABE_12, "AVICOLA,DEL CARMEN SA DE CV/", "0015", "CUENTA INEXISTENTE"),
        ])
        cheques, clabe = preregistro.leer_respuesta_preregistro(datos)
        self.assertEqual(cheques, {"banco": "0002", "tipo_cuenta": "01", "cuenta": "0" * 9 + SUCURSAL + CUENTA,
                                   "sucursal": SUCURSAL, "cuenta_banamex": CUENTA, "clabe": None,
                                   "beneficiario": "BRUNO RICARDO,HERNANDEZ/SILVA",
                                   "codigo": "0000", "mensaje": "ALTA APLICADA"})
        self.assertEqual(clabe, {"banco": "0072", "tipo_cuenta": "40", "cuenta": "00" + CLABE_12,
                                 "sucursal": None, "cuenta_banamex": None, "clabe": CLABE_12,
                                 "beneficiario": "AVICOLA,DEL CARMEN SA DE CV/",
                                 "codigo": "0015", "mensaje": "CUENTA INEXISTENTE"})

    def test_un_registro_corto_no_revienta(self):
        """Si la línea no llega a la posición del código (o del mensaje), esos dos quedan vacíos en
        lugar de tirar la carga entera."""
        corto = registro("0002", "01", "0" * 9 + SUCURSAL + CUENTA, "X,Y/Z", "0000", "ALTA APLICADA")[:150]
        (leido,) = preregistro.leer_respuesta_preregistro(archivo_respuesta([corto]))
        self.assertEqual((leido["codigo"], leido["mensaje"]), ("", ""))
        self.assertEqual(leido["cuenta_banamex"], CUENTA)

    def test_aplicar_registra_rechaza_y_es_idempotente(self):
        datos = archivo_respuesta([
            registro("0002", "01", "0" * 9 + SUCURSAL + CUENTA, "BRUNO RICARDO,HERNANDEZ/SILVA", "0000", "ALTA APLICADA"),
            registro("0072", "40", "00" + CLABE_12, "AVICOLA,DEL CARMEN SA DE CV/", "0015", "CUENTA INEXISTENTE"),
            registro("0072", "40", "00" + OTRA_CLABE_12, "QUIEN,SABE/", "0000", "ALTA APLICADA"),
        ])
        file_url = subir(datos)
        frappe.set_user(TESORERIA)
        r = api.aplicar_respuesta_preregistro(file_url)
        self.assertEqual((r["registradas"], r["rechazadas"]), (1, 1))
        self.assertEqual(r["sin_coincidencia"], ["00" + OTRA_CLABE_12])
        self.assertEqual(tuple(frappe.db.get_value("Bank Account", self.cta_06,
                                                   ["estado_preregistro", "preregistro_respuesta"])),
                         ("Registrada", "0000 ALTA APLICADA"))
        self.assertEqual(tuple(frappe.db.get_value("Bank Account", self.cta_12,
                                                   ["estado_preregistro", "preregistro_respuesta"])),
                         ("Rechazada", "0015 CUENTA INEXISTENTE"))
        # Cargar dos veces el mismo archivo no cambia nada ni cuenta doble.
        self.assertEqual(api.aplicar_respuesta_preregistro(file_url), r)
        self.assertEqual(frappe.db.get_value("Bank Account", self.cta_06, "estado_preregistro"), "Registrada")

    def test_aplicar_es_de_tesoreria(self):
        file_url = subir(archivo_respuesta([registro("0002", "01", "0" * 9 + SUCURSAL + CUENTA, "X,Y/Z", "0000", "ALTA APLICADA")]))
        frappe.set_user(REVISOR)
        with self.assertRaises(frappe.PermissionError):
            api.aplicar_respuesta_preregistro(file_url)

    @unittest.skipUnless(os.path.isfile(MUESTRA_REAL), f"no está la muestra real del banco ({MUESTRA_REAL})")
    def test_la_muestra_real_del_banco(self):
        """La única respuesta real que hay: 68 altas de cuentas Banamex de cheques, todas aplicadas.
        Confirma las posiciones del lector sobre datos del banco y no sobre un archivo que arma la
        propia prueba. NO se copia al repo ni se imprime su contenido: son datos de empleados."""
        with open(MUESTRA_REAL, "rb") as f:
            registros = preregistro.leer_respuesta_preregistro(f.read())
        self.assertEqual(len(registros), 68)
        self.assertEqual({r["codigo"] for r in registros}, {"0000"})
        self.assertEqual({r["mensaje"] for r in registros}, {"ALTA APLICADA"})
        self.assertEqual({r["banco"] for r in registros}, {"0002"})
        self.assertEqual({r["tipo_cuenta"] for r in registros}, {"01"})
        # 68 cuentas distintas: si el lector leyera otras posiciones saldría siempre la misma (los
        # doce dígitos que van antes de la cuenta son iguales en los 68 registros).
        self.assertEqual(len({(r["sucursal"], r["cuenta_banamex"]) for r in registros}), 68)
        self.assertTrue(all(re.fullmatch(r"\d{4}", r["sucursal"]) for r in registros))
        self.assertTrue(all(re.fullmatch(r"\d{7}", r["cuenta_banamex"]) for r in registros))
        # Todos los beneficiarios del banco vienen con la estructura NOMBRES,PATERNO/MATERNO.
        self.assertTrue(all(r["beneficiario"].count(",") == 1 and r["beneficiario"].count("/") == 1
                            for r in registros))

    # --- marcar a mano lo que ya estaba dado de alta ----------------------------------------------

    def test_marcar_registrada_es_de_tesoreria(self):
        frappe.set_user(REVISOR)
        with self.assertRaises(frappe.PermissionError):
            api.marcar_registrada(json.dumps([self.cta_12]))
        frappe.set_user(TESORERIA)
        self.assertEqual(api.marcar_registrada(json.dumps([self.cta_12, self.cta_06]))["registradas"], 2)
        for cuenta in (self.cta_12, self.cta_06):
            estado, respuesta = frappe.db.get_value("Bank Account", cuenta,
                                                    ["estado_preregistro", "preregistro_respuesta"])
            self.assertEqual(estado, "Registrada")
            self.assertIn("marcada a mano", respuesta)
            self.assertIn(TESORERIA, respuesta)

    def test_marcar_registrada_exige_cuenta_verificada(self):
        frappe.db.set_value("Bank Account", self.cta_12, "verificada", 0)
        frappe.set_user(TESORERIA)
        with self.assertRaises(frappe.ValidationError):
            api.marcar_registrada([self.cta_12])
        self.assertEqual(frappe.db.get_value("Bank Account", self.cta_12, "estado_preregistro"), "Sin registrar")

    # --- el candado del lote y el borrado del estado ----------------------------------------------

    def test_el_lote_exige_que_la_cuenta_este_pre_registrada(self):
        partidas = [{"factura": self.moral.name, "importe": 100}]
        with self.assertRaises(frappe.ValidationError):
            crear_lotes(pruebas_comun.EMPRESA, FECHA, partidas)
        # Con la cuenta registrada en el banco, el mismo lote sale.
        frappe.db.set_value("Bank Account", self.cta_12, "estado_preregistro", "Registrada")
        (lote,) = crear_lotes(pruebas_comun.EMPRESA, FECHA, partidas)
        self.assertTrue(lote)

    def test_sin_exigir_preregistro_el_lote_pasa(self):
        """El candado se puede apagar en la configuración: mientras Tesorería termine de dar de alta
        a los proveedores que ya existían, los pagos no pueden quedarse parados."""
        frappe.db.set_single_value("Configuracion CxP", "exigir_preregistro", 0)
        self.addCleanup(frappe.db.set_single_value, "Configuracion CxP", "exigir_preregistro", 1)
        (lote,) = crear_lotes(pruebas_comun.EMPRESA, FECHA, [{"factura": self.moral.name, "importe": 100}])
        self.assertEqual(frappe.db.get_value("Bank Account", self.cta_12, "estado_preregistro"), "Sin registrar")
        self.assertTrue(lote)

    def test_cambiar_la_clabe_regresa_la_cuenta_a_sin_registrar(self):
        """Es el mismo punto donde se pierde la verificación de Tesorería: con otra CLABE, el alta
        que autorizó el banco ya no es de esa cuenta."""
        frappe.db.set_value("Bank Account", self.cta_12, {"estado_preregistro": "Registrada",
                                                          "preregistro_respuesta": "0000 ALTA APLICADA"})
        doc = frappe.get_doc("Bank Account", self.cta_12)
        doc.clabe = OTRA_CLABE_12
        doc.save()
        doc.reload()
        self.assertEqual(doc.estado_preregistro, "Sin registrar")
        self.assertFalse(doc.preregistro_respuesta)
        self.assertFalse(doc.preregistro_enviado_el)
        self.assertFalse(doc.verificada)
